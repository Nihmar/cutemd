"""Hover popup that previews the target of a Markdown/wikilink."""

import csv
import zipfile
from pathlib import Path

from PySide6.QtCore import QEvent, QPoint, QSize, Qt, QThread, QTimer, Signal
from PySide6.QtGui import QFont, QPixmap, QTextOption
from PySide6.QtPdf import QPdfDocument
from PySide6.QtWidgets import (
    QFrame,
    QLabel,
    QSizePolicy,
    QTextBrowser,
    QTextEdit,
    QVBoxLayout,
    QWidget,
)

from markdown.document_renderers import (
    docx_to_html,
    epub_to_html,
    pptx_to_html,
    xlsx_to_html,
)
from core.constants import CSV_EXTS, IMG_EXTS, MD_EXTS, PDF_EXTS
from core.logging import setup_logging

_LOG = setup_logging("cutemd.link_preview_popup")
from ui.syntax_highlighter import MarkdownHighlighter


# ---------------------------------------------------------------------------
# Async preview worker — renders ALL formats in a background thread
# ---------------------------------------------------------------------------


class _PreviewRenderThread(QThread):
    """Renders any file format to preview content in a background thread.

    Emits ``result(str, str)`` where the first value is the preview content
    (HTML or plain text) and the second is a type hint:
    ``"html"`` or ``"text"``.
    """
    result = Signal(str, str)  # (content, content_type: "html"|"text")

    def __init__(self, path: Path, ext: str, css: str, parent=None):
        super().__init__(parent)
        self._path = path
        self._ext = ext
        self._css = css

    def run(self) -> None:
        try:
            if self._ext in (".xlsx",):
                html = xlsx_to_html(self._path, self._css)
                self.result.emit(html, "html")
            elif self._ext in (".docx",):
                html = docx_to_html(self._path, self._css)
                self.result.emit(html, "html")
            elif self._ext in (".pptx",):
                html = pptx_to_html(self._path, self._css)
                self.result.emit(html, "html")
            elif self._ext in (".epub",):
                html = epub_to_html(self._path, self._css)
                self.result.emit(html, "html")
            elif self._ext in (".csv", ".tsv"):
                text = self._render_csv()
                self.result.emit(text, "text")
            elif self._ext in (".pdf",):
                text = self._render_pdf()
                self.result.emit(text, "text")
            elif self._ext in (".cbz",):
                text = self._render_cbz()
                self.result.emit(text, "text")
            else:
                # Plain text / markdown — read directly
                try:
                    text = self._path.read_text(encoding="utf-8")
                    lines = text.split("\n")
                    if len(lines) > 500:
                        text = "\n".join(lines[:500]) + "\n\u2026 (truncated)"
                    self.result.emit(text, "text")
                except Exception:
                    self.result.emit(f"[Cannot read: {self._path.name}]", "text")
        except Exception as exc:
            self.result.emit(f"[Error rendering: {exc}]", "text")

    def _render_csv(self) -> str:
        import csv
        delimiter = "\t" if self._path.suffix.lower() == ".tsv" else ","
        with open(self._path, encoding="utf-8", errors="replace") as f:
            reader = csv.reader(f, delimiter=delimiter)
            rows = list(reader)[:100]
        if not rows:
            return ""
        col_widths = [0] * max(len(r) for r in rows)
        for row in rows:
            for i, cell in enumerate(row):
                if i < len(col_widths):
                    col_widths[i] = max(col_widths[i], len(cell))
        return "\n".join(
            " \u2502 ".join(c.ljust(col_widths[i]) for i, c in enumerate(row))
            for row in rows
        )

    def _render_pdf(self) -> str:
        return f"[PDF preview not available in popup: {self._path.name}]"

    def _render_cbz(self) -> str:
        import zipfile
        try:
            with zipfile.ZipFile(self._path) as zf:
                names = sorted(zf.namelist())
                img_names = [n for n in names if Path(n).suffix.lower()
                             in {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".svg", ".webp"}]
                return f"[CBZ: {len(img_names)} pages — {self._path.name}]"
        except Exception:
            return f"[Cannot read: {self._path.name}]"


# ---------------------------------------------------------------------------
# LinkPreviewPopup
# ---------------------------------------------------------------------------


class LinkPreviewPopup(QFrame):
# ---------------------------------------------------------------------------
    """Popup window that shows a live preview of a linked file.

    Appears near the mouse cursor when hovering over a link in the editor.
    """

    _MD_EXTS = MD_EXTS
    _IMG_EXTS = IMG_EXTS
    _PDF_EXTS = PDF_EXTS
    _CSV_EXTS = CSV_EXTS
    _DOCX_EXTS = frozenset({".docx"})
    _PPTX_EXTS = frozenset({".pptx"})
    _XLSX_EXTS = frozenset({".xlsx"})
    _CBZ_EXTS = frozenset({".cbz"})
    _EPUB_EXTS = frozenset({".epub"})

    _MAX_W = 520
    _MAX_H = 380

    def __init__(self, parent: QWidget | None = None) -> None:
        super().__init__(parent)
        self.setWindowFlags(
            Qt.WindowType.Popup
            | Qt.WindowType.FramelessWindowHint
        )
        self.setAttribute(Qt.WidgetAttribute.WA_ShowWithoutActivating)
        self.setAttribute(Qt.WidgetAttribute.WA_DeleteOnClose, False)
        self.setFrameStyle(QFrame.Shape.Box | QFrame.Shadow.Plain)

        self._path: Path | None = None
        self._hide_timer = QTimer(self)
        self._hide_timer.setSingleShot(True)
        self._hide_timer.setInterval(250)
        self._hide_timer.timeout.connect(self._maybe_hide)
        self._mouse_over = False

        # CBZ gallery state — stores only image names, data loaded on demand
        self._cbz_path: Path | None = None
        self._cbz_names: list[str] = []
        self._cbz_index: int = 0

        # PPTX slide state
        self._pptx_total: int = 0
        self._pptx_index: int = 0

        layout = QVBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(0)

        # --- Header bar with file path ---
        self._header = QLabel()
        self._header.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self._header.setWordWrap(False)
        self._header.setFixedHeight(22)
        self._header.setStyleSheet(
            "background: palette(dark); color: palette(bright-text); font-size: 10px; padding: 2px 6px;"
        )
        layout.addWidget(self._header)

        # --- Text / HTML preview ---
        self._editor = QTextBrowser()
        self._editor.setReadOnly(True)
        self._editor.setLineWrapMode(QTextEdit.LineWrapMode.NoWrap)
        self._editor.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        self._editor.setVerticalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAsNeeded)
        self._editor.setWordWrapMode(QTextOption.WrapMode.NoWrap)
        self._editor.setSizePolicy(
            QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding
        )
        self._highlighter: MarkdownHighlighter | None = None
        self._editor.installEventFilter(self)
        self._editor.viewport().installEventFilter(self)
        layout.addWidget(self._editor)

        # --- Image viewer ---
        self._image_label = QLabel()
        self._image_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self._image_label.setSizePolicy(
            QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding
        )
        self._image_label.setStyleSheet("background: palette(base);")
        self._image_label.hide()
        layout.addWidget(self._image_label)

        self.setMaximumSize(self._MAX_W, self._MAX_H)
        self.resize(self._MAX_W, self._MAX_H)

    # ------------------------------------------------------------------
    # Public API
    # ------------------------------------------------------------------

    def set_theme(self, theme: str) -> None:
        """Update the syntax highlighter theme."""
        if self._highlighter is not None:
            self._highlighter.set_theme(theme)

    def show_for_path(
        self, path: Path, screen_pos, editor_font: QFont | None = None
    ) -> None:
        """Load *path* and show the popup at *screen_pos*.

        All formats are rendered in a background thread so the popup
        appears instantly with a loading indicator. Images are the
        only exception — they load fast enough synchronously.
        """
        if path == self._path and self.isVisible():
            self._move_within_screen(screen_pos)
            return

        self._path = path
        self._hide_timer.stop()
        self._mouse_over = True

        # Cancel any pending render thread.
        if hasattr(self, "_render_thread") and self._render_thread.isRunning():
            self._render_thread.result.disconnect()
            self._render_thread.quit()
            self._render_thread.wait(500)

        was_visible = self.isVisible()
        if was_visible:
            self.hide()

        self._cbz_path = None
        self._cbz_names = []
        self._pptx_total = 0
        self._pptx_index = 0
        self._editor.viewport().setCursor(Qt.CursorShape.ArrowCursor)
        self._image_label.setCursor(Qt.CursorShape.ArrowCursor)

        self._header.setText(str(path))

        ext = path.suffix.lower()

        # Images — fast enough to keep synchronous.
        if ext in self._IMG_EXTS:
            self._show_image(path)
            self._move_within_screen(screen_pos)
            self.setVisible(True)
            self.raise_()
            return

        # Everything else — show popup immediately with loading,
        # render in background thread.
        self._image_label.hide()
        self._editor.show()
        self._editor.setLineWrapMode(QTextEdit.LineWrapMode.WidgetWidth)
        self._editor.setWordWrapMode(QTextOption.WrapMode.WordWrap)
        if self._highlighter is not None:
            self._highlighter.setDocument(None)
        self._editor.setPlainText(self.tr("Loading\u2026"))
        self._move_within_screen(screen_pos)
        self.setVisible(True)
        self.raise_()
        self._start_async_render(path, ext)

    def _start_async_render(self, path: Path, ext: str) -> None:
        """Launch background rendering for a heavy document format."""
        self._render_thread = _PreviewRenderThread(path, ext, "", self)
        self._render_thread.result.connect(self._on_async_render_done)
        self._render_thread.start()

    def _on_async_render_done(self, content: str, content_type: str) -> None:
        """Called when the background render finishes."""
        if content_type == "html":
            self._editor.setHtml(content)
        else:
            self._editor.setPlainText(content)
            # Apply monospace font for text/CSV previews.
            from PySide6.QtGui import QFont
            mono = QFont("Consolas", 9) if QFont("Consolas").exactMatch() else QFont("monospace", 9)
            self._editor.setFont(mono)
        self._editor.verticalScrollBar().setValue(0)

    def hide_popup(self) -> None:
        """Schedule hide after a short delay (to avoid flicker)."""
        self._mouse_over = False
        self._hide_timer.start()

    def cancel_hide(self) -> None:
        """Cancel pending hide (e.g., mouse entered the popup)."""
        self._hide_timer.stop()
        self._mouse_over = True

    # ------------------------------------------------------------------
    # Internal preview methods
    # ------------------------------------------------------------------

    _MAX_PREVIEW_LINES = 500

    def _show_text(self, path: Path, editor_font: QFont | None) -> None:
        self._image_label.hide()
        self._editor.show()
        self._editor.setLineWrapMode(QTextEdit.LineWrapMode.NoWrap)
        self._editor.setWordWrapMode(QTextOption.WrapMode.NoWrap)

        try:
            text = path.read_text(encoding="utf-8")
        except (OSError, UnicodeDecodeError):
            text = self.tr("[Cannot read file: {}]").format(path.name)

        # Truncate very large files to avoid excessive memory / rendering cost.
        lines = text.split("\n")
        if len(lines) > self._MAX_PREVIEW_LINES:
            text = "\n".join(lines[: self._MAX_PREVIEW_LINES]) + "\n… (truncated)"

        self._editor.setPlainText(text)

        if editor_font is not None:
            self._editor.setFont(editor_font)

        ext = path.suffix.lower()
        if ext in self._MD_EXTS:
            if self._highlighter is None:
                self._highlighter = MarkdownHighlighter(self._editor.document())
            else:
                self._highlighter.setDocument(self._editor.document())
        elif self._highlighter is not None:
            self._highlighter.setDocument(None)

        self._editor.verticalScrollBar().setValue(0)

    def _show_image(self, path: Path) -> None:
        self._editor.hide()
        if self._highlighter is not None:
            self._highlighter.setDocument(None)

        pixmap = QPixmap(str(path))
        if pixmap.isNull():
            self._image_label.setText(self.tr("[Cannot display: {}]").format(path.name))
        else:
            scaled = pixmap.scaled(
                self._MAX_W - 10,
                self._MAX_H - 30,
                Qt.AspectRatioMode.KeepAspectRatio,
                Qt.TransformationMode.SmoothTransformation,
            )
            self._image_label.setPixmap(scaled)

        self._image_label.show()

    def _show_pdf(self, path: Path) -> None:
        self._editor.hide()
        if self._highlighter is not None:
            self._highlighter.setDocument(None)

        doc = QPdfDocument()
        doc.load(str(path))
        if doc.pageCount() == 0:
            self._image_label.setText(self.tr("[Cannot render: {}]").format(path.name))
            self._image_label.show()
            return

        page_size = doc.pagePointSize(0)
        aspect = (
            page_size.height() / page_size.width() if page_size.width() > 0 else 1.4
        )
        w = min(self._MAX_W - 10, int(page_size.width()))
        h = min(self._MAX_H - 30, int(w * aspect))
        if h > self._MAX_H - 30:
            h = self._MAX_H - 30
            w = int(h / aspect) if aspect > 0 else 1

        img = doc.render(0, QSize(w, h))
        self._image_label.setPixmap(QPixmap.fromImage(img))
        self._image_label.show()

    def _show_csv(self, path: Path) -> None:
        self._image_label.hide()
        self._editor.show()
        self._editor.setLineWrapMode(QTextEdit.LineWrapMode.NoWrap)
        self._editor.setWordWrapMode(QTextOption.WrapMode.NoWrap)
        if self._highlighter is not None:
            self._highlighter.setDocument(None)

        try:
            delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
            with open(path, encoding="utf-8", errors="replace") as f:
                reader = csv.reader(f, delimiter=delimiter)
                rows = list(reader)[:100]
        except Exception:
            self._editor.setPlainText(self.tr("[Cannot read: {}]").format(path.name))
            return

        if not rows:
            self._editor.setPlainText("")
            return

        col_widths = [0] * max(len(r) for r in rows)
        for row in rows:
            for i, cell in enumerate(row):
                if i < len(col_widths):
                    col_widths[i] = max(col_widths[i], len(cell))

        lines = [" \u2502 ".join(c.ljust(col_widths[i]) for i, c in enumerate(row)) for row in rows]
        self._editor.setPlainText("\n".join(lines))
        mono = QFont("Consolas", 9) if QFont("Consolas").exactMatch() else QFont("monospace", 9)
        self._editor.setFont(mono)

    def _show_cbz(self, path: Path) -> None:
        self._editor.hide()
        if self._highlighter is not None:
            self._highlighter.setDocument(None)

        self._cbz_path = None
        self._cbz_names = []
        self._cbz_index = 0

        try:
            with zipfile.ZipFile(path) as zf:
                names = sorted(zf.namelist())
                self._cbz_names = [
                    n for n in names
                    if Path(n).suffix.lower() in self._IMG_EXTS
                ]
            self._cbz_path = path
        except (zipfile.BadZipFile, FileNotFoundError):
            self._image_label.setText(self.tr("[Cannot read: file is corrupted]"))
            self._image_label.show()
            self._editor.hide()
            return
        except Exception:
            pass

        if not self._cbz_names:
            self._image_label.setText(self.tr("[Cannot display: {}]").format(path.name))
            self._image_label.show()
            return

        self._show_cbz_page(0)

    def _show_cbz_page(self, index: int) -> None:
        """Show the *index*-th image from the CBZ archive (loaded on demand)."""
        name = self._cbz_names[index]
        try:
            with zipfile.ZipFile(self._cbz_path) as zf:
                data = zf.read(name)
        except Exception:
            self._image_label.setText(
                self.tr("[Cannot display page: {}]").format(Path(name).name)
            )
            self._image_label.show()
            return
        pixmap = QPixmap()
        if not pixmap.loadFromData(data) or pixmap.isNull():
            self._image_label.setText(
                self.tr("[Cannot display page: {}]").format(Path(name).name)
            )
            self._image_label.show()
            return

        total = len(self._cbz_names)
        self._header.setText(f"\u25c0  {self._path.name}  [{index + 1}/{total}]  \u25b6")
        self._image_label.setCursor(Qt.CursorShape.PointingHandCursor)
        scaled = pixmap.scaled(
            self._MAX_W - 10,
            self._MAX_H - 30,
            Qt.AspectRatioMode.KeepAspectRatio,
            Qt.TransformationMode.SmoothTransformation,
        )
        self._image_label.setPixmap(scaled)
        self._image_label.show()

    def mousePressEvent(self, event) -> None:
        """Navigate CBZ/PPTX pages with mouse clicks (left half=prev, right half=next)."""
        if event.button() == Qt.MouseButton.LeftButton:
            if self._cbz_names:
                w = self._image_label.width()
                if w > 0 and event.position().x() > w / 2:
                    self._cbz_index = (self._cbz_index + 1) % len(self._cbz_names)
                else:
                    self._cbz_index = (self._cbz_index - 1) % len(self._cbz_names)
                self._show_cbz_page(self._cbz_index)
                return
            if self._pptx_total > 0:
                w = self._editor.width()
                if w > 0 and event.position().x() > w / 2:
                    self._navigate_pptx(1)
                else:
                    self._navigate_pptx(-1)
                return
        super().mousePressEvent(event)

    def wheelEvent(self, event) -> None:
        """Navigate CBZ/PPTX pages with the mouse wheel."""
        if self._cbz_names:
            if event.angleDelta().y() < 0:
                self._cbz_index = (self._cbz_index + 1) % len(self._cbz_names)
                self._show_cbz_page(self._cbz_index)
            elif event.angleDelta().y() > 0:
                self._cbz_index = (self._cbz_index - 1) % len(self._cbz_names)
                self._show_cbz_page(self._cbz_index)
            event.accept()
            return
        if self._pptx_total > 0:
            if event.angleDelta().y() < 0:
                self._navigate_pptx(1)
            elif event.angleDelta().y() > 0:
                self._navigate_pptx(-1)
            event.accept()
            return
        super().wheelEvent(event)

    def keyPressEvent(self, event) -> None:
        """Navigate CBZ/PPTX pages with Left/Right arrow keys."""
        if self._cbz_names:
            if event.key() == Qt.Key.Key_Right:
                self._cbz_index = (self._cbz_index + 1) % len(self._cbz_names)
                self._show_cbz_page(self._cbz_index)
            elif event.key() == Qt.Key.Key_Left:
                self._cbz_index = (self._cbz_index - 1) % len(self._cbz_names)
                self._show_cbz_page(self._cbz_index)
            return
        if self._pptx_total > 0:
            if event.key() == Qt.Key.Key_Right:
                self._navigate_pptx(1)
            elif event.key() == Qt.Key.Key_Left:
                self._navigate_pptx(-1)
            return
        super().keyPressEvent(event)

    def eventFilter(self, obj, event) -> bool:
        """Intercept events from the editor (and its viewport) for CBZ/PPTX navigation."""
        if obj is self._editor or obj is self._editor.viewport():
            if event.type() == QEvent.Type.MouseButtonPress:
                if self._pptx_total > 0:
                    _LOG.debug("eventFilter MousePress: pptx_total=%d", self._pptx_total)
                    self.mousePressEvent(event)
                    return True
            elif event.type() == QEvent.Type.Wheel:
                _LOG.debug("eventFilter Wheel: cbz=%d pptx_total=%d",
                            len(self._cbz_names), self._pptx_total)
                if self._cbz_names or self._pptx_total > 0:
                    self.wheelEvent(event)
                    return True
            elif event.type() == QEvent.Type.KeyPress:
                key = event.key()
                _LOG.debug("eventFilter KeyPress: key=%d cbz=%d pptx_total=%d",
                            key, len(self._cbz_names), self._pptx_total)
                if self._cbz_names and key in (Qt.Key.Key_Left, Qt.Key.Key_Right):
                    self.keyPressEvent(event)
                    return True
                if self._pptx_total > 0 and key in (Qt.Key.Key_Left, Qt.Key.Key_Right):
                    self.keyPressEvent(event)
                    return True
        return super().eventFilter(obj, event)

    def _show_epub(self, path: Path) -> None:
        self._image_label.hide()
        self._editor.show()
        if self._highlighter is not None:
            self._highlighter.setDocument(None)

        self._editor.setLineWrapMode(QTextEdit.LineWrapMode.WidgetWidth)
        self._editor.setWordWrapMode(QTextOption.WrapMode.WordWrap)
        html = epub_to_html(path, "")
        self._editor.setHtml(html)
        self._editor.verticalScrollBar().setValue(0)

    def _show_xlsx(self, path: Path) -> None:
        self._image_label.hide()
        self._editor.show()
        self._editor.setLineWrapMode(QTextEdit.LineWrapMode.NoWrap)
        self._editor.setWordWrapMode(QTextOption.WrapMode.WordWrap)
        if self._highlighter is not None:
            self._highlighter.setDocument(None)

        try:
            html = xlsx_to_html(path, "")
        except ImportError:
            html = f"<p>{self.tr('[Package openpyxl required for .xlsx preview]')}</p>"
        except Exception:
            html = f"<p>{self.tr('[Cannot read: {}]').format(path.name)}</p>"

        self._editor.setHtml(html)
        self._editor.verticalScrollBar().setValue(0)

    def _show_docx(self, path: Path) -> None:
        self._image_label.hide()
        self._editor.show()
        self._editor.setLineWrapMode(QTextEdit.LineWrapMode.NoWrap)
        self._editor.setWordWrapMode(QTextOption.WrapMode.WordWrap)
        if self._highlighter is not None:
            self._highlighter.setDocument(None)

        try:
            html = docx_to_html(path, "")
        except ImportError:
            html = f"<p>{self.tr('[Package python-docx required for .docx preview]')}</p>"
        except Exception:
            html = f"<p>{self.tr('[Cannot read: {}]').format(path.name)}</p>"

        self._editor.setHtml(html)
        self._editor.verticalScrollBar().setValue(0)

    def _show_pptx(self, path: Path) -> None:
        self._image_label.hide()
        self._editor.show()
        self._editor.setLineWrapMode(QTextEdit.LineWrapMode.NoWrap)
        self._editor.setWordWrapMode(QTextOption.WrapMode.WordWrap)
        self._editor.viewport().setCursor(Qt.CursorShape.PointingHandCursor)
        if self._highlighter is not None:
            self._highlighter.setDocument(None)

        try:
            from pptx import Presentation
            prs = Presentation(str(path))
            self._pptx_total = len(list(prs.slides))
            self._pptx_index = 0
            _LOG.debug("_show_pptx: %s total=%d", path.name, self._pptx_total)
            html = pptx_to_html(path, "", slide_index=0)
            self._update_pptx_header()
        except ImportError:
            _LOG.debug("_show_pptx: ImportError python-pptx")
            html = f"<p>{self.tr('[Package python-pptx required for .pptx preview]')}</p>"
            self._pptx_total = 0
        except Exception as exc:
            _LOG.debug("_show_pptx: %s", exc)
            html = f"<p>{self.tr('[Cannot read: {}]').format(path.name)}</p>"
            self._pptx_total = 0

        self._editor.setHtml(html)
        self._editor.verticalScrollBar().setValue(0)

    def _update_pptx_header(self) -> None:
        if self._pptx_total > 0:
            self._header.setText(
                f"\u25c0  {self._path.name}  [{self._pptx_index + 1}/{self._pptx_total}]  \u25b6"
            )
            _LOG.debug("_update_pptx_header: index=%d total=%d",
                        self._pptx_index, self._pptx_total)

    def _navigate_pptx(self, delta: int) -> None:
        if self._pptx_total <= 0:
            _LOG.debug("_navigate_pptx: no slides")
            return
        new_index = (self._pptx_index + delta) % self._pptx_total
        if new_index == self._pptx_index:
            return
        self._pptx_index = new_index
        _LOG.debug("_navigate_pptx: delta=%d new_index=%d total=%d",
                    delta, self._pptx_index, self._pptx_total)
        try:
            html = pptx_to_html(self._path, "", slide_index=self._pptx_index)
            self._editor.setHtml(html)
            self._editor.verticalScrollBar().setValue(0)
            self._update_pptx_header()
        except Exception as exc:
            _LOG.debug("_navigate_pptx error: %s", exc)

    # ------------------------------------------------------------------
    # Delayed hide
    # ------------------------------------------------------------------

    def _maybe_hide(self) -> None:
        if not self._mouse_over:
            self.hide()

    def _move_within_screen(self, pos) -> None:
        """Move to *pos* (screen coords), nudging the popup back on-screen if needed."""
        from PySide6.QtGui import QScreen
        from PySide6.QtWidgets import QApplication

        app = QApplication.instance()
        if app is None:
            self.move(pos)
            return

        x = pos.x()
        y = pos.y()
        sz = self.size()

        screen: QScreen | None = app.screenAt(QPoint(x, y))
        if screen is None:
            screen = app.primaryScreen()
        if screen is None:
            self.move(pos)
            return

        geo = screen.availableGeometry()

        if x + sz.width() > geo.right():
            x = geo.right() - sz.width()
        if x < geo.left():
            x = geo.left()
        if y + sz.height() > geo.bottom():
            y = geo.bottom() - sz.height()
        if y < geo.top():
            y = geo.top()

        _LOG.debug("_move_within_screen: pos=%s final=(%d,%d) sz=%s geo=%s",
                    pos, x, y, sz, (geo.x(), geo.y(), geo.width(), geo.height()))

        # Popup windows use screen coordinates for move().
        self.move(x, y)
        _LOG.debug("_move_within_screen: actual pos=%s actual global=%s",
                    self.pos(), self.mapToGlobal(QPoint(0, 0)))

    # ------------------------------------------------------------------
    def enterEvent(self, event) -> None:
        super().enterEvent(event)
        self.cancel_hide()

    def leaveEvent(self, event) -> None:
        super().leaveEvent(event)
        self.hide_popup()
